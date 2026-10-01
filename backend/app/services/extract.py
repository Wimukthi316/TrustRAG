"""Pull plain text out of an uploaded document, for the "attach a file" control.

Deliberately free of any heavy import -- no torch, no transformers, no FastAPI --
so it can be unit tested on its own and the route in main.py stays thin. It takes
bytes and a filename and returns text; the model is only ever given that text,
never the file. This is plumbing for the paste box, so it lives outside
schema.py and carries no novelty claim.

The only third-party dependencies are pypdf (PDF) and python-docx (.docx); plain
text and markdown need neither, and both are imported lazily so a checkout
missing them still serves text files.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Optional

# One source of truth for what the control accepts. The frontend copies these
# values into its `accept=` attribute and its wording; there is no import across
# the language boundary, so if one moves the other must move in the same commit.
MAX_BYTES = 10 * 1024 * 1024  # 10 MB -- a RAG source document is text, not a scan
MAX_PAGES = 100  # defensive cap; a RAGTruth context is a page or two
SUPPORTED = (".pdf", ".docx", ".txt", ".md")


class ExtractError(ValueError):
    """The file could not be turned into text. The message is shown to the user."""


@dataclass
class Extracted:
    text: str
    pages: Optional[int]  # PDF page count, else None
    note: Optional[str]  # a caveat worth showing, e.g. "this looks like a scan"


def _suffix(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot != -1 else ""


def _clean(text: str) -> str:
    """Normalise whitespace so pasted-in text does not arrive ragged.

    PDF extraction in particular leaves trailing spaces and runs of blank lines
    that would otherwise show up in the document box.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    out: list[str] = []
    blank = False
    for line in text.split("\n"):
        line = line.rstrip()
        if line:
            out.append(line)
            blank = False
        elif not blank:
            out.append("")
            blank = True
    return "\n".join(out).strip()


def extract_text(filename: str, data: bytes) -> Extracted:
    """Turn the bytes of one uploaded file into clean text.

    Raises ExtractError with a message meant for the person, never a traceback.
    """
    if not data:
        raise ExtractError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise ExtractError(
            f"The file is {len(data) / 1_048_576:.1f} MB; the limit is "
            f"{MAX_BYTES // 1_048_576} MB. A source document should be text, "
            "not a scanned image."
        )
    suffix = _suffix(filename)
    if suffix == ".pdf":
        return _from_pdf(data)
    if suffix == ".docx":
        return _from_docx(data)
    if suffix in (".txt", ".md"):
        return _from_text(data)
    raise ExtractError(
        f"{suffix or 'That file type'} is not supported. Attach a PDF, a Word "
        ".docx, or a .txt file."
    )


def _from_pdf(data: bytes) -> Extracted:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as exc:  # pragma: no cover
        raise ExtractError(
            "PDF support needs the 'pypdf' package (pip install pypdf)."
        ) from exc
    try:
        reader = PdfReader(BytesIO(data))
    except PdfReadError as exc:
        raise ExtractError(f"This PDF could not be read: {exc}.") from exc
    if reader.is_encrypted:
        # An empty-password decrypt covers the common owner-locked-but-readable
        # case; a real password we cannot guess.
        try:
            reader.decrypt("")
        except Exception as exc:  # noqa: BLE001
            raise ExtractError(
                "This PDF is password-protected. Remove the password and try again."
            ) from exc
    pages = len(reader.pages)
    if pages > MAX_PAGES:
        raise ExtractError(
            f"This PDF has {pages} pages; the limit is {MAX_PAGES}. Paste the "
            "relevant passage instead."
        )
    chunks: list[str] = []
    for page in reader.pages:
        try:
            chunks.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            # One unreadable page should not sink the rest of the document.
            chunks.append("")
    text = _clean("\n\n".join(chunks))
    note = None
    if not text:
        note = (
            "No text came out of this PDF. It is most likely a scan or an image, "
            "which would need OCR -- this tool reads text, it does not do OCR."
        )
    return Extracted(text=text, pages=pages, note=note)


def _from_docx(data: bytes) -> Extracted:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise ExtractError(
            "Word support needs the 'python-docx' package (pip install python-docx)."
        ) from exc
    try:
        document = docx.Document(BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractError(
            f"This does not look like a valid .docx file ({exc}). If it is an "
            "older .doc, open it in Word and save it as .docx."
        ) from exc
    parts = [p.text for p in document.paragraphs]
    # Tables carry text too, and a source table is exactly the kind of thing a
    # data-to-text answer is checked against.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = _clean("\n".join(parts))
    if not text:
        raise ExtractError("No text was found in this document.")
    return Extracted(text=text, pages=None, note=None)


def _from_text(data: bytes) -> Extracted:
    text = None
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ExtractError("This text file is not in a recognised encoding.")
    text = _clean(text)
    if not text:
        raise ExtractError("The file is empty.")
    return Extracted(text=text, pages=None, note=None)
